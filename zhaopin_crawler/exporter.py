# -*- coding: utf-8 -*-
"""
exporter.py —— 导出 Excel / CSV
================================
生成一个多工作表的 Excel：
    采集说明 / 岗位明细 / 分类统计 / 二级分类 / 城市统计 /
    薪资分布 / 技能热度 / 公司招聘量 / 经验与学历
其中分类统计和城市统计带柱状图，打开就能看。
"""

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="2F5597")
HEADER_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="微软雅黑", size=14, bold=True, color="1F3864")
NOTE_FONT = Font(name="微软雅黑", size=10, color="C00000")
BODY_FONT = Font(name="微软雅黑", size=10)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

# 岗位明细表：列名 -> (宽度, 字段名)
DETAIL_COLUMNS = [
    ("序号", 6, None),
    ("岗位名称", 28, "job_name"),
    ("一级分类", 12, "category_l1"),
    ("二级分类", 16, "category_l2"),
    ("薪资", 15, "salary_desc"),
    ("月薪下限K", 11, "salary_min"),
    ("月薪上限K", 11, "salary_max"),
    ("月薪中值K", 11, "salary_mid"),
    ("年薪估算(万)", 13, "salary_annual"),
    ("薪资档位", 12, "salary_band"),
    ("城市", 10, "city"),
    ("区域", 12, "area"),
    ("街道", 14, "street"),
    ("经验要求", 11, "experience"),
    ("经验档位", 11, "exp_level"),
    ("学历要求", 10, "degree"),
    ("公司名称", 32, "company"),
    ("公司性质", 11, "company_type"),
    ("行业", 20, "industry"),
    ("公司规模", 13, "scale"),
    ("技能标签", 36, "skills"),
    ("福利待遇", 26, "welfare"),
    ("岗位链接", 50, "job_link"),
    ("是否实习", 9, "is_intern"),
    ("发布时间", 20, "publish_time"),
    ("采集时间", 19, "crawl_time"),
    ("搜索关键词", 12, "keyword"),
    ("岗位编号", 26, "job_no"),
]


def _value(record: dict, field: str):
    if field is None:
        return None
    v = record.get(field)
    if isinstance(v, (list, tuple)):
        return " / ".join(str(x) for x in v if x)
    return "" if v is None else v


def _write_table(ws, headers, widths, rows, start_row=1, freeze=True, as_number=None):
    """写一张带样式的表。as_number: 需要按数字格式显示的列索引集合（1-based）。"""
    as_number = as_number or set()

    for c, (name, width) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=start_row, column=c, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(c)].width = width

    for r, row in enumerate(rows, start_row + 1):
        for c, val in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = CENTER if c in as_number or isinstance(val, (int, float)) else LEFT
            if c in as_number and val not in (None, ""):
                cell.number_format = "0.0"

    if freeze:
        ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    if rows:
        ws.auto_filter.ref = (
            f"A{start_row}:{get_column_letter(len(headers))}{start_row + len(rows)}"
        )
    ws.row_dimensions[start_row].height = 26


def _add_bar_chart(ws, title, data_first_row, data_last_row, cat_col=1, val_col=2,
                   anchor="E2", y_title="岗位数"):
    chart = BarChart()
    chart.type = "bar"
    chart.title = title
    chart.y_axis.title = y_title
    chart.x_axis.title = None
    chart.height = 9
    chart.width = 18
    data = Reference(ws, min_col=val_col, min_row=data_first_row - 1,
                     max_row=data_last_row)
    cats = Reference(ws, min_col=cat_col, min_row=data_first_row, max_row=data_last_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.legend = None
    ws.add_chart(chart, anchor)


def _sheet_note(ws, title, stats, meta):
    """采集说明页。"""
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 80

    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:B1")

    is_demo = meta.get("is_demo", False)
    row = 3
    if is_demo:
        ws.cell(row=row, column=1, value="【演示数据】")
        ws.cell(row=row, column=2,
                value="本文件由 --demo 模式生成，岗位内容为示例数据，"
                      "用于验证分类与导出流程，不是真实在招岗位。")
        ws.cell(row=row, column=1).font = NOTE_FONT
        ws.cell(row=row, column=2).font = NOTE_FONT
        row += 2

    o = stats["overall"]
    rows = [
        ("生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("数据来源", meta.get("source", "智联招聘（sou.zhaopin.com）")),
        ("采集模式", meta.get("mode", "")),
        ("采集参数", meta.get("params", "")),
        ("岗位总数", o["岗位总数"]),
        ("有效薪资样本", o["有效薪资样本"]),
        ("平均月薪", f"{o['平均月薪K']} K/月" if o["平均月薪K"] else "—"),
        ("月薪中位数", f"{o['月薪中位数K']} K/月" if o["月薪中位数K"] else "—"),
        ("月薪区间", f"{o['月薪最低K']}K ~ {o['月薪最高K']}K" if o["平均月薪K"] else "—"),
        ("采集耗时", f"{meta.get('elapsed', 0)} 秒"),
        ("", ""),
        ("数据说明", "真实采集时：智联的检索页每个「关键词×城市」组合固定返回 20 条且不翻页，"
                    "岗位总量靠多组合累加，重复岗位已按「岗位编号」去重；"
                    "本文件若标注「演示数据」则不是真实岗位。"),
        ("分类口径", "按岗位名称关键词匹配（先具体后宽泛），名称匹配不上时用技能标签兜底，"
                    "仍未命中的归入「其他」。规则见 classifier.py。"),
        ("薪资口径", "优先取智联返回的 salaryReal 数字区间（元/月）折算为「千元/月」；"
                    "「14薪」等按实际月数估算年薪；金额过小（如 150-300 元/次）标为「按次计薪」，"
                    "不并入平均月薪，避免拉偏统计。"),
        ("使用提示", "「岗位明细」表已开启筛选，可按分类/城市/薪资档位自由筛选；"
                    "「分类统计」等表右侧附图表。"),
        ("合规提示", "本工具仅用于个人学习与求职分析，请控制采集频率、遵守网站 robots 协议与用户协议，"
                    "不要用于商业转售或大规模骚扰性抓取。"),
    ]
    for k, v in rows:
        row += 1
        kc = ws.cell(row=row, column=1, value=k)
        vc = ws.cell(row=row, column=2, value=v)
        kc.font = Font(name="微软雅黑", size=10, bold=True)
        kc.alignment = Alignment(vertical="top")
        vc.font = BODY_FONT
        vc.alignment = Alignment(vertical="top", wrap_text=True)
        if k:
            ws.row_dimensions[row].height = 32 if len(str(v)) > 40 else 18


def _save_with_fallback(save_fn, out_path: Path, logger=None):
    """
    保存文件。如果目标文件正被 Excel 打开（Windows 会锁定），
    自动改存成「原名_时间戳.后缀」，而不是直接报错退出。
    """
    try:
        save_fn(out_path)
        return out_path
    except PermissionError:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        alt = out_path.with_name(f"{out_path.stem}_{ts}{out_path.suffix}")
        save_fn(alt)
        if logger:
            logger.warning(
                f"{out_path.name} 正被其他程序占用（可能正在 Excel 里打开），"
                f"已改存为 {alt.name}"
            )
        return alt


def export_excel(records, stats, meta, out_path: Path, logger=None):
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()

    # 1. 采集说明
    ws = wb.active
    ws.title = "采集说明"
    _sheet_note(ws, "智联招聘岗位采集报告", stats, meta)

    # 2. 岗位明细
    ws = wb.create_sheet("岗位明细")
    headers = [c[0] for c in DETAIL_COLUMNS]
    widths = [c[1] for c in DETAIL_COLUMNS]
    number_cols = {6, 7, 8, 9}
    body = []
    for i, rec in enumerate(records, 1):
        body.append([i] + [_value(rec, c[2]) for c in DETAIL_COLUMNS[1:]])
    _write_table(ws, headers, widths, body, as_number=number_cols)
    # 第 1 列序号宽度小一点，链接列点击可跳转
    link_col = headers.index("岗位链接") + 1
    for r in range(2, len(body) + 2):
        cell = ws.cell(row=r, column=link_col)
        if isinstance(cell.value, str) and cell.value.startswith("http"):
            cell.hyperlink = cell.value
            cell.font = Font(name="微软雅黑", size=9, color="0563C1", underline="single")

    # 3. 分类统计
    ws = wb.create_sheet("分类统计")
    rows = [[r["分类"], r["岗位数"], r["占比%"], r["平均月薪K"]] for r in stats["category"]]
    _write_table(ws, ["一级分类", "岗位数", "占比%", "平均月薪K"], [16, 10, 10, 12], rows,
                 as_number={2, 3, 4})
    if rows:
        _add_bar_chart(ws, "各一级分类岗位数量", 2, len(rows) + 1, anchor="G2")

    # 4. 二级分类
    ws = wb.create_sheet("二级分类")
    rows = [[r["一级分类"], r["二级分类"], r["岗位数"], r["平均月薪K"]] for r in stats["sub_category"]]
    _write_table(ws, ["一级分类", "二级分类", "岗位数", "平均月薪K"], [16, 20, 10, 12], rows,
                 as_number={3, 4})

    # 5. 城市统计
    ws = wb.create_sheet("城市统计")
    rows = [[r["城市"], r["岗位数"], r["平均月薪K"], r["最高月薪K"]] for r in stats["city"]]
    _write_table(ws, ["城市", "岗位数", "平均月薪K", "最高月薪K"], [14, 10, 12, 12], rows,
                 as_number={2, 3, 4})
    if rows:
        _add_bar_chart(ws, "各城市岗位数量", 2, len(rows) + 1, anchor="G2")

    # 6. 薪资分布
    ws = wb.create_sheet("薪资分布")
    rows = [[r["薪资档位"], r["岗位数"]] for r in stats["salary_dist"]]
    _write_table(ws, ["薪资档位", "岗位数"], [16, 10], rows, as_number={2})

    # 7. 技能热度
    ws = wb.create_sheet("技能热度")
    rows = [[r["技能标签"], r["出现岗位数"]] for r in stats["skills"]]
    _write_table(ws, ["技能标签", "出现岗位数"], [26, 12], rows, as_number={2})
    if rows:
        _add_bar_chart(ws, "技能被提及次数 Top 15", 2, min(len(rows), 15) + 1, anchor="E2")

    # 8. 公司招聘量
    ws = wb.create_sheet("公司招聘量")
    rows = [[r["公司"], r["在招岗位数"], r["平均月薪K"]] for r in stats["companies"]]
    _write_table(ws, ["公司", "在招岗位数", "平均月薪K"], [36, 12, 12], rows, as_number={2, 3})

    # 9. 经验与学历
    ws = wb.create_sheet("经验与学历")
    ws["A1"] = "经验要求分布"
    ws["A1"].font = Font(name="微软雅黑", size=11, bold=True)
    rows = [[r["项目"], r["岗位数"]] for r in stats["experience"]]
    _write_table(ws, ["经验档位", "岗位数"], [16, 10], rows, start_row=2, freeze=False, as_number={2})

    start = len(rows) + 5
    ws.cell(row=start, column=1, value="学历要求分布").font = Font(name="微软雅黑", size=11, bold=True)
    rows2 = [[r["项目"], r["岗位数"]] for r in stats["degree"]]
    _write_table(ws, ["学历档位", "岗位数"], [16, 10], rows2, start_row=start + 1, freeze=False, as_number={2})

    out_path = _save_with_fallback(wb.save, out_path, logger)
    if logger:
        logger.info(f"Excel 已生成：{out_path.name}")
    return out_path


def export_csv(records, out_path: Path, logger=None):
    import csv

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    headers = [c[0] for c in DETAIL_COLUMNS]

    def write(target: Path):
        with open(target, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow(headers)
            for i, rec in enumerate(records, 1):
                writer.writerow([i] + [_value(rec, c[2]) for c in DETAIL_COLUMNS[1:]])

    out_path = _save_with_fallback(write, out_path, logger)
    if logger:
        logger.info(f"CSV 已生成：{out_path.name}")
    return out_path
